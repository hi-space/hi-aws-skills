#!/usr/bin/env python3
"""Convert reference files in <work>/sources/ into plain text under <work>/sources/text/.

    python3 ingest_sources.py <work> [--force]

Handles: .pdf (pdftotext, fallback pypdf), .docx (python-docx or pandoc), .pptx (python-pptx),
.md/.txt/.json/.yaml/.csv (copied), .drawio (node labels and edge labels extracted from the XML,
brief.md kept as-is). Video/audio files are listed as "not readable"; ask the author for a transcript.
Prints a registry table you can paste into 00-brief.md.
"""
from __future__ import annotations

import argparse
import html
import re
import shutil
import subprocess
import sys
import xml.etree.ElementTree as ET
from pathlib import Path

TEXT_EXT = {".md", ".txt", ".json", ".yaml", ".yml", ".csv", ".log", ".py", ".ts", ".tf", ".hcl", ".url"}
MEDIA_EXT = {".mp4", ".mov", ".mkv", ".mp3", ".wav", ".m4a", ".webm"}


def pdf_to_text(src: Path, dst: Path) -> str:
    if shutil.which("pdftotext"):
        r = subprocess.run(["pdftotext", "-layout", str(src), str(dst)], capture_output=True, text=True)
        if r.returncode == 0 and dst.exists() and dst.stat().st_size > 0:
            return "pdftotext"
    try:
        import pypdf  # type: ignore

        reader = pypdf.PdfReader(str(src))
        pages = []
        for i, page in enumerate(reader.pages, 1):
            pages.append(f"\n\n===== page {i} =====\n")
            pages.append(page.extract_text() or "")
        dst.write_text("".join(pages), encoding="utf-8")
        return "pypdf"
    except Exception as exc:  # noqa: BLE001
        return f"failed: {exc}"


def docx_to_text(src: Path, dst: Path) -> str:
    try:
        import docx  # type: ignore

        d = docx.Document(str(src))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        dst.write_text("\n".join(parts), encoding="utf-8")
        return "python-docx"
    except Exception:  # noqa: BLE001
        pass
    if shutil.which("pandoc"):
        r = subprocess.run(["pandoc", str(src), "-t", "gfm", "-o", str(dst)], capture_output=True, text=True)
        if r.returncode == 0:
            return "pandoc"
        return f"failed: {r.stderr.strip()[:200]}"
    return "failed: no python-docx or pandoc"


def pptx_to_text(src: Path, dst: Path) -> str:
    try:
        from pptx import Presentation  # type: ignore
    except Exception as exc:  # noqa: BLE001
        return f"failed: python-pptx missing ({exc})"
    prs = Presentation(str(src))
    out = []
    for i, slide in enumerate(prs.slides, 1):
        out.append(f"\n===== slide {i} =====")
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    t = "".join(r.text for r in para.runs).strip()
                    if t:
                        out.append(t)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    out.append(" | ".join(c.text.strip() for c in row.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                out.append(f"[notes] {notes}")
    dst.write_text("\n".join(out), encoding="utf-8")
    return "python-pptx"


def drawio_to_text(src: Path, dst: Path) -> str:
    try:
        root = ET.fromstring(src.read_text(encoding="utf-8"))
    except ET.ParseError as exc:
        return f"failed: {exc}"
    cells = {}
    for c in root.iter("mxCell"):
        cells[c.get("id")] = c
    lines = [f"# nodes and edges extracted from {src.name}", "", "## nodes"]
    for cid, c in cells.items():
        v = c.get("value") or ""
        v = re.sub(r"<[^>]+>", " ", html.unescape(v)).strip()
        if c.get("vertex") == "1" and v:
            lines.append(f"- {cid}: {v}")
    lines.append("")
    lines.append("## edges (source -> target: label)")
    for cid, c in cells.items():
        if c.get("edge") == "1":
            s = cells.get(c.get("source"))
            t = cells.get(c.get("target"))
            sv = re.sub(r"<[^>]+>", " ", html.unescape(s.get("value") or "")).strip() if s is not None else "?"
            tv = re.sub(r"<[^>]+>", " ", html.unescape(t.get("value") or "")).strip() if t is not None else "?"
            lab = re.sub(r"<[^>]+>", " ", html.unescape(c.get("value") or "")).strip()
            lines.append(f"- {sv} -> {tv}: {lab}")
    dst.write_text("\n".join(lines), encoding="utf-8")
    return "drawio-xml"


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("work")
    ap.add_argument("--force", action="store_true", help="re-convert even if text exists")
    args = ap.parse_args()
    work = Path(args.work)
    src_dir = work / "sources"
    out_dir = src_dir / "text"
    if not src_dir.is_dir():
        print(f"no sources dir at {src_dir}", file=sys.stderr)
        return 2
    out_dir.mkdir(exist_ok=True)

    rows = []
    n = 0
    for f in sorted(p for p in src_dir.iterdir() if p.is_file() and p.name != "README.md"):
        if f.name.endswith(".digest.md"):
            continue
        n += 1
        sid = f"S{n:02d}"
        ext = f.suffix.lower()
        dst = out_dir / (f.stem + ".txt")
        if ext in MEDIA_EXT:
            rows.append((sid, f.name, "video/audio", "not readable", "ask for transcript or written summary"))
            continue
        if dst.exists() and not args.force:
            rows.append((sid, f.name, ext.lstrip("."), "converted", f"text/{dst.name} (kept)"))
            continue
        if ext == ".pdf":
            how = pdf_to_text(f, dst)
        elif ext == ".docx":
            how = docx_to_text(f, dst)
        elif ext == ".pptx":
            how = pptx_to_text(f, dst)
        elif ext == ".drawio":
            how = drawio_to_text(f, dst)
        elif ext in TEXT_EXT:
            shutil.copyfile(f, dst)
            how = "copied"
        else:
            rows.append((sid, f.name, ext.lstrip("."), "unknown", "read manually or convert by hand"))
            continue
        status = "failed" if how.startswith("failed") else "converted"
        rows.append((sid, f.name, ext.lstrip("."), status, f"text/{dst.name} via {how}" if status == "converted" else how))

    print("| ID | 파일/링크 | 종류 | 읽기 가능 | 비고 |")
    print("|---|---|---|---|---|")
    for r in rows:
        print("| " + " | ".join(r) + " |")
    failed = [r for r in rows if r[3] in ("failed", "not readable", "unknown")]
    if failed:
        print(f"\n{len(failed)} source(s) need attention (see above).")
    return 0


if __name__ == "__main__":
    sys.exit(main())
